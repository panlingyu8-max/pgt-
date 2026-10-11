"""Markdown -> 实验方案 docx（python-docx；宋体/Times New Roman 12pt，1.5倍行距，黑色标题）。
用法: python3 build.py input.md output.docx figures_dir
"""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src, dst, figdir = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(sec, side, Cm(2.54))
TEXT_W = 21.0 - 2 * 2.54

def set_font(obj, size=None, bold=None, east='宋体'):
    obj.font.name = 'Times New Roman'
    if size: obj.font.size = Pt(size)
    if bold is not None: obj.font.bold = bold
    obj.font.color.rgb = RGBColor(0, 0, 0)
    rpr = obj.element.get_or_add_rPr() if hasattr(obj, 'element') else obj._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for k in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        rf.attrib.pop(qn('w:' + k), None)
    rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:eastAsia'), east)

normal = doc.styles['Normal']
set_font(normal, 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(3)
title_st = doc.styles['Title']; set_font(title_st, 18, True, '黑体')
title_st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_st.paragraph_format.space_after = Pt(12)
for lvl, size in ((1, 15), (2, 14), (3, 12.5), (4, 12)):
    st = doc.styles[f'Heading {lvl}']; set_font(st, size, True, '黑体')
    st.paragraph_format.space_before = Pt(10 if lvl < 3 else 6)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.3
for el in list(doc.styles.element.iter()):
    if el.tag == qn('w:pBdr'): el.getparent().remove(el)

PLACE = re.compile(r'(【[^】]*】)')
def add_runs(p, text, size=None, bold=False):
    parts = PLACE.split(text)
    for part in parts:
        if not part: continue
        # **bold** support
        for j, seg in enumerate(re.split(r'\*\*(.+?)\*\*', part)):
            if not seg: continue
            r = p.add_run(seg)
            set_font(r, size, bold or j % 2 == 1)
            if PLACE.fullmatch(part):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW

def cell_borders(cell, header):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '808080'); b.append(e)
    tcpr.append(b)
    mar = OxmlElement('w:tcMar')
    for edge, v in (('top', '25'), ('bottom', '25'), ('left', '90'), ('right', '90')):
        e = OxmlElement('w:' + edge); e.set(qn('w:w'), v); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcpr.append(mar)
    if header:
        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'E8EAED'); tcpr.append(sh)

def col_widths(rows):
    n = len(rows[0])
    lens = [0.5*max(len(r[j]) if j < len(r) else 0 for r in rows)+0.5*sum(len(r[j]) if j < len(r) else 0 for r in rows)/len(rows) for j in range(n)]
    lens = [min(max(l, 5), 40) for l in lens]
    tot = sum(lens)
    return [TEXT_W * l / tot for l in lens]

def add_table(rows):
    n = len(rows[0]); widths = col_widths(rows)
    if rows[0][0].startswith(('消融参数','电生理')):
        widths = [5.0] + [(TEXT_W-5.0)/(n-1)]*(n-1)
    if rows[0][0].startswith(('档位比较点','瓣膜与冠脉','右室体部参照点')):
        widths = [6.6] + [(TEXT_W-6.6)/(n-1)]*(n-1)
    if rows[0][0].startswith('指标（单位）'):
        widths = [4.4] + [(TEXT_W-4.4)/(n-1)]*(n-1)
    t = doc.add_table(rows=0, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j,col in enumerate(t.columns): col.width = Cm(widths[j])
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(n):
            c = cells[j]; c.width = Cm(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_borders(c, i == 0)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
            add_runs(p, row[j] if j < len(row) else '', 10, i == 0)
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        if rows[0][0].startswith(('消融参数','电生理','术中项目','消融能量参数','指标（单位）','档位比较点','瓣膜与冠脉','右室体部参照点')):
            h=OxmlElement('w:trHeight'); h.set(qn('w:val'),'310'); h.set(qn('w:hRule'),'atLeast'); trpr.append(h)
    grid = t._tbl.tblGrid
    for j,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[j]*567)))
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = Pt(6)

def caption(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    add_runs(p, text, 10.5, True)

# header / footer
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_runs(hp, '流出道PFA与RF猪实验方案 V3.1', 9)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    r = fp.add_run(); set_font(r, 9)
    if kind:
        f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), kind); r._r.append(f)
    else:
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r._r.append(it)

lines = src.read_text().splitlines(); i = 0
in_refs = False
while i < len(lines):
    line = lines[i].rstrip(); s = line.strip()
    if not s: i += 1; continue
    if s.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', c) for c in row): rows.append(row)
            i += 1
        add_table(rows); continue
    m = re.match(r'\[\[FIGURE:([^|]+)\|(.+)\]\]', s)
    if m:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(figdir / m.group(1)), width=Cm(14.5))
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(c, m.group(2), 10.5, True)
        i += 1; continue
    if s.startswith('# '):
        doc.add_paragraph(s[2:], style='Title')
    elif s.startswith('## '):
        h = s[3:]
        in_refs = (h == '参考文献')
        hp = doc.add_paragraph(h, style='Heading 1')
        if h.startswith('附表'): hp.paragraph_format.page_break_before = True
    elif s.startswith('### '): doc.add_paragraph(s[4:], style='Heading 2')
    elif s.startswith('#### '): doc.add_paragraph(s[5:], style='Heading 3')
    elif re.match(r'^表\d+　', s):
        caption(s)
    elif s.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); add_runs(p, s[2:])
        p.paragraph_format.space_after = Pt(1)
    elif re.match(r'^\d+\. ', s):
        p = doc.add_paragraph()
        if in_refs:
            p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.75)
            add_runs(p, s, 10)
        else:
            p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, s)
    else:
        p = doc.add_paragraph(); add_runs(p, s)
        if s.startswith(('项目名称', '申请类别', '关键词')):
            p.paragraph_format.line_spacing = 1.3
    i += 1

for st_name in ('List Bullet',):
    set_font(doc.styles[st_name], 12)
doc.core_properties.title = src.read_text().splitlines()[0].lstrip('# ').strip()
doc.core_properties.author = ''
for el in list(doc.styles.element.iter())+list(doc.element.body.iter()):
    if el.tag in (qn('w:keepNext'), qn('w:keepLines')):
        el.getparent().remove(el)
doc.save(dst)
print(dst)
