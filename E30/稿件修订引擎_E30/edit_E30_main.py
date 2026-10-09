# -*- coding: utf-8 -*-
"""E30 main text (base: E29 track-changes file with all revisions accepted).
1. Single study population (3,682 -> 131 uniformly excluded -> 3,551) with two analysis sets: baseline phenotype set 3,045 and
   paired response set 844 procedures (837 patients); 372 patients in both. 5,397/5,680, 2,425, 'verified/confirmed pre-procedural',
   'separately assembled / not nested', and report-availability weighting removed. 'cohort' -> 'set' for the two analysis sets.
2. Baseline numbers replaced by the 3,045-patient values (R48/R50/R55); S1B values 627 -> 601, <=0.9 -> <=0.5 (R56).
3. New post hoc joint-improvement factor analysis (R57; Supplementary Table S13, Figure 4B); Figure 4C BMI only, RA/LVEDd curves
   moved to Supplementary Figure S2.
4. Introduction aims, Discussion, Limitations, What's new, Conclusion, and abstract rewritten along the burden -> joint improvement ->
   who improves jointly -> 1-month state and later course line.
Figure images (1, 2, 4) are NOT swapped here: the redrawn files live in E28_图像与绘图依据/ on the author's computer.
Usage: python edit_E30_main.py <accepted_E29.docx> <out_trackchanges.docx>"""
import sys, re, os
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import tc
from runlevel import tracked_sub
tc.AUTHOR = 'Claude (E30 修订)'
tc.DATE = '2026-10-09T12:00:00Z'
BASE, OUT = sys.argv[1:3]
d = Document(BASE); body = d.element.body
TXT = lambda p: tc.accepted_text(p._p)
P = lambda el: Paragraph(el, d._body)
def allp(): return [P(p) for p in body.iter(qn('w:p'))]
def hits(prefix, exact=False): return [p for p in allp() if (TXT(p).strip() == prefix if exact else TXT(p).startswith(prefix))]
def one(prefix, exact=False):
    h = hits(prefix, exact); assert len(h) == 1, (prefix, len(h)); return h[0]
def rev(p, pairs):
    t = TXT(p)
    for old, new in pairs:
        assert old in t, (old, t[:120]); t = t.replace(old, new)
    tc.revise(p, t)
def wc_main(doc):
    T = [p.text for p in doc.paragraphs]
    i0 = next(i for i, t in enumerate(T) if t.strip() == 'Introduction'); i1 = next(i for i, t in enumerate(T) if t.strip() == 'Funding')
    ts = [t for t in T[i0:i1] if not re.match(r'^(Figure|Table) \d\.', t) and not t.startswith('Baseline analysis grade ≥2 defines')]
    return sum(len(re.findall(r'\S+', t)) for t in ts)
def wc_abs(doc):
    T = [p.text for p in doc.paragraphs]
    i0 = next(i for i, t in enumerate(T) if t.strip() == 'Abstract'); i1 = next(i for i, t in enumerate(T) if t.startswith('Key Words'))
    return sum(len(re.findall(r'\S+', t)) for t in T[i0 + 1:i1] if t.strip() not in ('Background and Aims', 'Methods', 'Results', 'Conclusion'))
def cell_par(tr, i): return P(tr.findall(qn('w:tc'))[i].find(qn('w:p')))
def set_cell(tr, i, new):
    p = cell_par(tr, i)
    if TXT(p) != new: tc.revise(p, new)
WC0 = wc_main(d)

# ================================================================ abstract
tc.revise(one('We characterized baseline phenotypes in 2,322 patients'),
    'Among 3,551 AF ablation patients with an in-hospital baseline echocardiogram, baseline phenotypes were characterized in 3,045 '
    'and 1-month responses in 844 procedures (837 patients; 343 paroxysmal, 501 persistent AF) with baseline MR or TR. '
    'Improvement was a ≥1-analysis-grade reduction. Twelve prespecified baseline factors were mutually adjusted and, post hoc, related to joint improvement.')
tc.revise(one('Baseline MR and/or TR was present in 54.1%'),
    'Baseline MR and/or TR was present in 51.6%; combined disease was more frequent in persistent than paroxysmal AF (37.4% vs 11.6%). '
    'At 1 month, MR improved in 47.2% and TR in 52.7% of valve series; among moderate-or-greater series, 43.3% of MR and 48.9% of TR regressed to mild or less. '
    'In combined disease, both valves improved in 34.4%, exceeding independent expectation in both AF types (observed/expected ratio, 1.36); '
    'blinded re-reading supported the association (odds ratio [OR], 4.01; 95% CI, 2.10–7.63). '
    'Joint improvement was more likely with higher body mass index (OR, 1.79; 1.33–2.40) and less likely with larger left atrium (OR, 0.62; 0.43–0.89). '
    'Joint improvement at 1 year was present in 83% after early joint improvement versus 19% after none.')
tc.revise(one('Within 1 month after ablation, about half of MR and TR improved'),
    'About half of AF ablation candidates had functional MR or TR; one-fifth had both. '
    'After ablation, both valves tended to improve together, more often with higher body mass index and smaller atria, and early joint improvement usually persisted, '
    'supporting valve-by-valve reassessment at 1 month.')
# What's new: clinical information
for prefix, new in (
    ('• Parallel analyses in paroxysmal and persistent AF identified excess joint',
     '• About half of patients undergoing AF ablation had functional MR or TR, and one in five had both; combined disease was three times as frequent in persistent as in paroxysmal AF.'),
    ('• Among study-defined moderate-or-greater series (analysis grade ≥2), 43% of MR',
     '• Within 1 month after ablation, MR and TR tended to improve together beyond chance in both AF types, and nearly half of moderate-or-greater regurgitation regressed to mild or less.'),
    ('• Discordant improvement occurred in 34% and 31%',
     '• Joint improvement was more likely with higher body mass index and smaller atria; with a larger left atrium, improvement was more often confined to one valve.'),
    ('• All isolated and combined valve series contributed to the analyses.',
     '• Joint improvement at 1 month usually persisted at 1 year (83% vs 19% when neither valve improved), supporting valve-by-valve echocardiographic reassessment at 1 month.')):
    tc.revise(one(prefix), new)

# ================================================================ introduction
tc.revise(one('We examined early MR/TR responses in parallel in paroxysmal and persistent AF.'),
    'We therefore addressed four questions in patients undergoing AF ablation, in parallel in paroxysmal and persistent AF: how common isolated and combined MR/TR are before ablation; '
    'whether both valves improve together within patients beyond chance; which baseline characteristics distinguish patients whose improvement extends to both valves; '
    'and whether the joint state at 1 month indicates the later course. Blinded re-reading and an aortic regurgitation (AR) negative control assessed robustness.')

# ================================================================ methods
tc.revise(one('Study populations and parallel analysis framework', exact=True), 'Study population and analysis sets')
tc.revise(one('This single-center retrospective study at West China Hospital'),
    'This single-center retrospective study at West China Hospital, Sichuan University, included adults (age ≥18 years) undergoing AF ablation between January 2020 and September 2024 '
    'whose index admission included an in-hospital baseline transthoracic echocardiography (TTE) report (n=3,682). After uniform exclusion of prior or concurrent valve surgery (n=74), '
    'organic valve disease (n=30), congenital heart disease (n=19), and obstructive hypertrophic cardiomyopathy (n=8), the study population comprised 3,551 patients (Figure 1). '
    'Two analysis sets addressed complementary questions: a baseline phenotype set of 3,045 patients characterized regurgitation burden and coexistence, and a paired response set of '
    '844 procedures in 837 patients assessed 1-month improvement; 372 patients belonged to both sets. Paroxysmal and persistent AF were analyzed separately within each set using common definitions.')
tc.revise(one('For the baseline phenotype cohort, screening began with 2,425 patients'),
    'In the baseline phenotype set, MR and TR grades were extracted from the TTE report of the index admission (weighted κ against manual abstraction, 0.96–0.97), and patients were '
    'classified as having neither valve, MR only, TR only, or both affected (analysis grade ≥1; Supplementary Table S1A). Registry grades were recorded only for patients with '
    'regurgitation and were not used to estimate prevalence.')
tc.revise(one('For the paired response cohort, we identified consecutive patients'),
    'For the paired response set, the study registry identified 1,445 admissions in 1,419 patients of the study population with MR or TR of analysis grade ≥1 (very mild or worse) on the '
    'in-hospital baseline TTE. Admissions without a 1-month TTE (n=601) were excluded and compared with included procedures (Supplementary Table S1B). The set comprised 844 procedures '
    'in 837 patients: isolated MR (n=132), isolated TR (n=325), and combined MR/TR (n=387). Seven patients underwent two ablations, with each procedure analyzed separately. '
    'For the 372 patients in both sets, paired-analysis baseline grades were used in both.')
for p in hits('Table 2. Baseline clinical, echocardiographic, and procedural characteristics of the paired response cohort'):
    rev(p, [('of the paired response cohort', 'of the paired response set')])
rev(one('The paroxysmal AF response cohort comprised 343'), [('The paroxysmal AF response cohort comprised', 'The paroxysmal AF group of the paired response set comprised')])
rev(one('The persistent AF response cohort comprised 501'), [('The persistent AF response cohort comprised', 'The persistent AF group comprised')])
rev(one('In the paroxysmal AF cohort, both valves improved'), [('In the paroxysmal AF cohort,', 'In paroxysmal AF,')])
rev(one('In the persistent AF cohort, both valves improved'), [('In the persistent AF cohort,', 'In persistent AF,')])
rev(one('The cohorts showed no detectable heterogeneity'), [('The cohorts showed', 'The AF-type groups showed')])
rev(one('The analyses addressed linked questions'),
    [('conditional joint-response models, TRV change,', 'conditional joint-response models, baseline factors of joint improvement, TRV change,')])
rev(one('In combined disease, MR–TR improvement associations were tested'),
    [('adjusted for baseline grades and AF type.',
      'adjusted for baseline grades and AF type. Also post hoc, the 12 baseline factors, with both baseline grades in place of phenotype, were related to improvement of both valves, '
      'to the number of improved valves (proportional-odds model), and, among procedures with at least one improved valve, to extension of improvement to both valves (Supplementary Table S13).')])
rev(one('In the baseline phenotype cohort, coexistence was compared'),
    [('In the baseline phenotype cohort,', 'In the baseline phenotype set,'),
     ('adjusted for age, sex, and AF type; and selection by report availability was assessed by inverse probability weighting.', 'adjusted for age, sex, and AF type.')])

# ================================================================ results: baseline
tc.revise(one('Among 2,322 patients undergoing a first AF ablation'),
    'Among 3,045 patients of the baseline phenotype set, 1,571 (51.6%) had MR and/or TR of analysis grade ≥1. Isolated MR occurred in 303 (10.0%), isolated TR in 605 (19.9%), '
    'and combined MR/TR in 663 (21.8%; 95% CI, 20.3–23.3; Figure 2A–B). Paroxysmal and persistent AF accounted for 1,848 and 1,197 patients, respectively (Table 1). '
    'These proportions describe ablation patients with an in-hospital baseline TTE report (Supplementary Table S1A).')
rev(one('In the paroxysmal AF population (n=1,367)'),
    [('In the paroxysmal AF population (n=1,367), 139 patients (10.2%) had isolated MR, 237 (17.3%) isolated TR, and 172 (12.6%)',
      'In paroxysmal AF (n=1,848), 188 patients (10.2%) had isolated MR, 291 (15.7%) isolated TR, and 215 (11.6%)'),
     ('(OR, 2.58; 95% CI, 1.83–3.64)', '(OR, 2.91; 95% CI, 2.14–3.96)')])
rev(one('In the persistent AF population (n=955)'),
    [('In the persistent AF population (n=955), 86 patients (9.0%) had isolated MR, 261 (27.3%) isolated TR, and 361 (37.8%)',
      'In persistent AF (n=1,197), 115 patients (9.6%) had isolated MR, 314 (26.2%) isolated TR, and 448 (37.4%)'),
     ('(OR, 2.87; 95% CI, 1.93–4.29)', '(OR, 2.83; 95% CI, 2.00–4.01)'),
     ('(prevalence ratio, 3.12; 95% CI, 2.67–3.65)', '(prevalence ratio, 3.32; 95% CI, 2.88–3.82)')])
rev(one('Overall, MR–TR co-occurrence exceeded the expectation'),
    [('(observed/expected ratio, 1.26; 95% CI, 1.20–1.30)', '(observed/expected ratio, 1.27; 95% CI, 1.22–1.32)'),
     ('grade ≥2 MR occurred in 38.1%, versus 26.2% with isolated MR; corresponding TR proportions were 53.7% and 29.9%',
      'grade ≥2 MR occurred in 37.7%, versus 22.8% with isolated MR; corresponding TR proportions were 52.8% and 27.9%')])

# ---------------------------------------------------------------- Table 1 (baseline phenotype set: R55_B)
t1title = hits('Table 1. Baseline characteristics by AF type')
for p in t1title:
    rev(p, [('in the baseline phenotype and paired response cohorts.', 'in the baseline phenotype and paired response analysis sets.')])
t1 = t1title[0]._p.getnext(); assert t1.tag == qn('w:tbl')
NEW = {'Baseline phenotype cohort (patients)': ['Baseline phenotype set (patients)'],
       'Patients, n': ['Patients, n', '1,848', '1,197', ''],
       'Age, years': ['Age, years', '63 (54, 71)', '65 (56, 71)', '0.017'],
       'Sex, female': ['Sex, female', '840 (45.5%)', '407 (34.0%)', '<0.001'],
       'Neither': ['Neither', '1,154 (62.4%)', '320 (26.7%)', ''],
       'MR only': ['MR only', '188 (10.2%)', '115 (9.6%)', ''],
       'TR only': ['TR only', '291 (15.7%)', '314 (26.2%)', ''],
       'MR and TR': ['MR and TR', '215 (11.6%)', '448 (37.4%)', ''],
       'MR analysis grade ≥2': ['MR analysis grade ≥2', '120 (6.5%)', '199 (16.6%)', '<0.001'],
       'TR analysis grade ≥2': ['TR analysis grade ≥2', '154 (8.3%)', '365 (30.5%)', '<0.001'],
       'LAD, mm': ['LAD, mm', '37.0 (34.0, 41.0)', '42.0 (39.0, 46.0)', '<0.001'],
       'LVEDd, mm': ['LVEDd, mm', '47.0 (45.0, 50.0)', '48.0 (45.0, 51.2)', '<0.001'],
       'RAD, mm': ['RAD, mm', '35.0 (32.0, 38.0)', '44.0 (39.0, 47.0)', '<0.001'],
       'RVD, mm': ['RVD, mm', '22.0 (20.0, 23.0)', '23.0 (21.0, 24.0)', '<0.001'],
       'LVEF, %': ['LVEF, %', '66 (62, 70)', '62 (56, 67)', '<0.001'],
       'Concomitant LAAO': ['Concomitant LAAO', '197 (10.7%)', '216 (18.0%)', '<0.001']}
done = set()
for tr in t1.iter(qn('w:tr')):
    k = TXT(cell_par(tr, 0)).strip()
    if k == 'Paired response cohort (procedures)':
        set_cell(tr, 0, 'Paired response set (procedures)'); break
    if k in NEW:
        for i, v in enumerate(NEW[k]): set_cell(tr, i, v)
        done.add(k)
assert done == set(NEW), set(NEW) - done

# ---------------------------------------------------------------- results: paired set
tc.revise(one('Paired response population', exact=True), 'Paired response analysis set')
rev(one('The separately assembled paired response cohort comprised 844'),
    [('The separately assembled paired response cohort comprised', 'The paired response set comprised'),
     ('The 627 eligible admissions without a 1-month TTE', 'The 601 eligible admissions without a 1-month TTE'),
     ('changed improvement rates by ≤0.9 percentage points', 'changed improvement rates by ≤0.5 percentage points')])

# ---------------------------------------------------------------- results: baseline factors -> joint improvement paragraph
gp = one('Among grade ≥2 series, the BMI association persisted')
rev(gp, [('Standardized probabilities of improvement across BMI and RA and LV end-diastolic diameters are shown in Figure 4B–D.',
          'Standardized probabilities of improvement are shown across BMI in Figure 4C and across RA and LV end-diastolic diameters in Supplementary Figure S2.')])
tc.insert_paragraph_after(gp,
    'In the 387 combined-disease procedures, higher BMI (26 vs 22 kg/m²) was associated with improvement of both valves (OR, 1.79; 95% CI, 1.33–2.40; q=0.001) and, '
    'among the 257 procedures with at least one improved valve, with extension of improvement to both valves (OR, 1.63; 95% CI, 1.13–2.34). '
    'Larger LA diameter (47 vs 40 mm) was associated with less joint improvement (OR, 0.62; 95% CI, 0.43–0.89) and less extension to both valves (OR, 0.54; 95% CI, 0.34–0.85), '
    'and larger RA diameter (48 vs 40 mm) with less joint improvement (OR, 0.69; 95% CI, 0.48–0.98); these atrial estimates did not reach q<0.05. '
    'The other factors showed no clear association (Figure 4B; Supplementary Table S13).')
for p in hits('Figure 4. Baseline factors associated with 1-month MR and TR improvement.'):
    tc.revise(p, 'Figure 4. Baseline factors associated with 1-month MR and TR improvement and with improvement of both valves.')

# ================================================================ discussion
tc.revise(one('Parallel analyses in paroxysmal and persistent AF identified excess joint MR/TR improvement alongside'),
    'In patients undergoing AF ablation, about half had functional MR or TR before the procedure and one in five had both. After ablation, the two valves tended to improve together '
    'beyond chance in both paroxysmal and persistent AF, consistent with a shared atrial substrate. Patients with higher BMI and smaller atria were more likely to improve in both valves, '
    'whereas with a larger LA, improvement was more often confined to one valve. The joint state at 1 month anticipated the course at about 1 year.')
tc.revise(one('Shared response patterns across AF types', exact=True), 'Regurgitation burden before ablation')
tc.revise(one('The 2,322-patient baseline phenotype cohort described'),
    'Among 3,045 patients with an in-hospital baseline TTE report, 51.6% had MR and/or TR and 21.8% had both. Combined regurgitation was three times as frequent in persistent as in '
    'paroxysmal AF and accompanied greater biatrial enlargement and higher grades. MR–TR co-occurrence remained after adjustment for chamber dimensions in both AF types, consistent with '
    'an atrial substrate.1,2,7,8 These data extend the combined phenotype described before cardioversion to a larger ablation population that includes both AF types.9 '
    'Individual-valve improvement was not confined to combined disease, although nonsignificant phenotype comparisons do not establish equivalence.')
tc.revise(one('Joint improvement and valve-specific variation', exact=True), 'Joint improvement and a shared atrial substrate')
tc.revise(one('Excess joint improvement in both AF types supports'),
    'Joint improvement exceeded chance in both AF types, without detectable heterogeneity between them, and blinded re-reading supported it overall. The two valves shared their response '
    'with atrial reverse remodeling: across tertiles of biatrial reduction, joint improvement rose from 15% to 56% and the proportion with neither valve improved fell, whereas discordance '
    'remained at about one-third. RA reduction accompanied TR improvement, and reduction of both atria accompanied MR improvement, in both AF types, extending earlier individual-valve '
    'findings.6,10,11 These patterns are consistent with a common atrial substrate; changes in atrial contraction and annular geometry are plausible contributors19–21 but were not directly '
    'measured, atrial size did not account for all coupling, and concurrent changes cannot establish which change preceded or caused another. SR accompanied TR improvement mainly in '
    'persistent AF, whereas SR was nearly universal in paroxysmal AF, limiting the rhythm contrast.')
tc.revise(one('Blinded re-reading supported the overall finding and the persistent AF estimate'),
    'Blinded re-reading supported the overall finding and the persistent AF estimate; the smaller paroxysmal estimate remained inconclusive. Adjustment for baseline severity reduced the '
    'possibility that the coupling reflected baseline grade alone. Shared acquisition conditions, rhythm-related grading variation, and regression to the mean remain possible contributors.')
tc.delete_paragraph(one('Shared atrial associations and differences in rhythm associations', exact=True)._p)
tc.delete_paragraph(one('The parallel framework also revealed differences')._p)
tc.revise(one('Baseline features and improvement expectations', exact=True), 'Who improves in both valves')
tc.revise(one('The 12 baseline factors connect the response pattern'),
    'Higher BMI was associated with improvement of each valve and, in combined disease, with improvement of both; when one valve improved, it was also associated with extension of '
    'improvement to the other. Conversely, a larger LA was associated with less joint improvement and, when one valve improved, with improvement confined to that valve; a larger RA showed '
    'a similar but weaker association. A larger atrium may reflect more advanced atrial remodeling with less reversible annular dilatation; neither this nor a load-related explanation for '
    'the BMI association was tested directly. These associations describe the baseline phenotype and do not imply that increasing weight improves regurgitation, and unindexed chamber '
    'diameters and residual body-size-related confounding require consideration. The joint-improvement analysis was post hoc, and its atrial estimates did not reach the false discovery threshold.')
tc.revise(one('The valve-specific associations refine these expectations.'),
    'The valve-specific associations refine these expectations. Larger RA diameter was associated with less TR improvement after mutual adjustment, while larger LV end-diastolic diameter '
    'was associated with more improvement; these conditional estimates may reflect different structural and loading backgrounds.22–25 Female sex was associated with less MR improvement, and '
    'higher baseline TR grade with more grade reduction, partly reflecting greater opportunity to decrease. Limited discrimination restricts individual prediction without negating the '
    'group-level associations.')
tc.revise(one('Persistence and implications for valve assessment', exact=True), 'The 1-month joint state and clinical implications')
tc.revise(one('Selected later imaging showed sustained and delayed improvement'),
    'In combined disease, early joint improvement persisted at about 1 year in 83% of procedures with late imaging, whereas only 19% of those with neither valve improved at 1 month showed '
    'joint improvement later (OR, 22.7). Loss or delay of TR improvement accompanied changes in RA size, consistent with previous observations of TR evolution.26 Late imaging was '
    'clinically indicated and available in about 30%, and the paroxysmal estimate rested on only 10 procedures; these findings cannot establish long-term durability across the full population.')
tc.revise(one('Current guidelines incorporate treatment of underlying AF'),
    'Current guidelines incorporate treatment of underlying AF and appropriate rhythm control into functional-regurgitation management, alongside assessment of valve-intervention '
    'indications.27 Our findings support valve-by-valve TTE assessment about 1 month after ablation. When both valves have improved, improvement usually persists. When only one or neither '
    'valve has improved, particularly in patients with a large atrium or lower BMI, targeted reassessment and further evaluation of the valve may be warranted, guided by residual severity, '
    'symptoms, and cardiac consequences. Whether this approach changes treatment sequencing or outcomes requires prospective evaluation.')
rev(one('Semi-quantitative clinical reports lacked comprehensive core-laboratory'),
    [('43% of eligible admissions lacked 1-month TTE, and the baseline cohort represented 43% of first ablations.',
      'patients without an in-hospital baseline TTE report were not included, and 42% of eligible admissions lacked 1-month TTE.'),
     ('higher-grade analyses were also post hoc.', 'higher-grade and joint-improvement factor analyses were also post hoc.')])
tc.revise(one('Paroxysmal and persistent AF shared excess joint MR/TR improvement'),
    'About half of patients undergoing AF ablation had functional MR or TR, and one in five had both. After ablation, the two valves tended to improve together beyond chance in both '
    'paroxysmal and persistent AF, more often in patients with higher BMI and smaller atria, and joint improvement at 1 month usually persisted at 1 year. Valve-specific reassessment at '
    '1 month can identify patients whose improvement is incomplete.')

# ================================================================ supplementary list, legends
rev(one('Tables S1A, S1B, S2, S3A–S3D, and S4–S12; Figure S1'), [('S4–S12; Figure S1', 'S4–S13; Figures S1 and S2')])
for p in hits('Figure 1. Study populations.'):
    tc.revise(p, 'Figure 1. Study population and analysis sets.')
tc.revise(one('Flow of the baseline phenotype cohort, starting from 2,425 patients'),
    'Flow of the study population: patients undergoing AF ablation between January 2020 and September 2024 whose index admission included an in-hospital baseline TTE report. '
    'Exclusions were applied uniformly before the analysis sets were defined. Organic valve disease comprised rheumatic heart disease, mitral stenosis, aortic stenosis of '
    'mild-to-moderate or worse, and valve prolapse; congenital heart disease comprised unrepaired atrial septal defect with shunt, residual shunt after repair or closure, and complex '
    'congenital heart disease. The baseline phenotype set comprised patients whose baseline TTE report was graded to estimate prevalence; the paired response set comprised procedures '
    'with MR and/or TR of analysis grade ≥1 (very mild or worse) on the baseline TTE and a 1-month TTE. Overall, 372 patients belonged to both sets, for whom the baseline grades of the '
    'paired analysis were used. Each set was analyzed separately in paroxysmal and persistent AF. Eligible admissions without a 1-month TTE were compared with included procedures in '
    'Supplementary Table S1B. AF indicates atrial fibrillation; HCM, hypertrophic cardiomyopathy; MR, mitral regurgitation; TR, tricuspid regurgitation; TTE, transthoracic echocardiography.')
rev(one('(A) MR/TR phenotypes (analysis grade ≥1) in the baseline phenotype cohort'),
    [('Panels A–B describe patients of the baseline phenotype cohort and panels C–F valve series or procedures of the paired response cohort; because the two populations were assembled separately and only partially overlap, no flow is drawn between them.',
      'Panels A–B describe patients of the baseline phenotype set (n=3,045) and panels C–F valve series or procedures of the paired response set; the two sets share 372 patients.')])
rev(one('(A) Mutually adjusted odds ratios (95% CIs) for ≥1-grade MR (all series)'),
    [('(B–D) Standardized probabilities (95% CIs) of ≥1-grade improvement at 1 month derived by g-computation from the model in A, across BMI for MR and TR (B), and across RA (C) and LV end-diastolic (D) diameters for TR.',
      '(B) Mutually adjusted odds ratios (95% CIs) of the same 12 factors, with both baseline grades in place of phenotype, for improvement of both valves (joint; n=387) and, among '
      'procedures with at least one improved valve, for extension of improvement to both valves (extension; n=257) in combined MR/TR; continuous factors are compared at the 75th '
      'versus 25th percentile in combined MR/TR (post hoc; Supplementary Table S13). (C) Standardized probabilities (95% CIs) of ≥1-grade improvement at 1 month derived by '
      'g-computation from the model in A across BMI for MR and TR; curves across RA and LV end-diastolic diameters for TR are shown in Supplementary Figure S2.')])
rev(one('Data are median (interquartile range) or n (%) of observed values. The baseline phenotype cohort counts'),
    [('The baseline phenotype cohort counts patients (first AF ablation admission) and the paired response cohort counts procedures; the two cohorts were assembled separately and shared 357 patients. In the baseline phenotype cohort, chamber dimensions were available from the same report in 2,095 (LAD), 2,289 (LVEDd), 2,098 (RAD), 2,212 (RVD), and 2,295 (LVEF) patients;',
      'The baseline phenotype set counts patients (index ablation admission) and the paired response set counts procedures; the two sets shared 372 patients. In the baseline phenotype set, chamber dimensions were available from the same report in 2,755 (LAD), 3,009 (LVEDd), 2,755 (RAD), 2,916 (RVD), and 3,009 (LVEF) patients;'),
     ('were abstracted only for the paired response cohort,', 'were abstracted only for the paired response set,')])

# ================================================================ global wording (last: tracked_sub leaves earlier revisions intact)
N = 0
for p in allp():
    for a, b in (('baseline phenotype cohort', 'baseline phenotype set'), ('Baseline phenotype cohort', 'Baseline phenotype set'),
                 ('paired response cohort', 'paired response set'), ('Paired response cohort', 'Paired response set')):
        N += tracked_sub(p._p, a, b)
print('global cohort->set substitutions:', N)

d.save(OUT)
from accept_all import accept_docx
tmp = OUT.replace('.docx', '_tmp.docx'); accept_docx(OUT, tmp)
D1 = Document(tmp); WC1, AB1 = wc_main(D1), wc_abs(D1); os.remove(tmp)
d = Document(OUT); body = d.element.body
wcp = one('Word count:')
m = re.search(r'Word count: ([\d,]+) main text; (\d+) abstract', TXT(wcp))
old_main = int(m.group(1).replace(',', ''))
rev(wcp, [(m.group(0), f'Word count: {old_main + WC1 - WC0:,} main text; {AB1} abstract')])
d.save(OUT)
print('saved', OUT, '| main', old_main, '->', old_main + WC1 - WC0, '| abstract', AB1)
