# -*- coding: utf-8 -*-
"""74. (E30) 统一研究人群流程（住院期间有基线 TTE 报告的消融患者）与 Table S1B 重算。
研究人群 = 电子档案中首次消融住院有 TTE 报告者（63 号脚本分析集来源，排除前）∪ 登记中 MR/TR ≥1 级的消融住院（1,479 次）的患者；
统一排除（瓣膜手术、器质性瓣膜病、先心病、梗阻性 HCM）；配对队列 844 次中的患者经团队评估纳入。
S1B：符合条件的登记住院中，属于被统一标准排除患者的住院一并去除后，重算纳入（844）与无 1 个月 TTE 者的比较及随访加权改善率。
输出：11_顶刊修订分析/tables/R56_A_unified_flow.csv、R56_B_exclusion_categories.csv、R56_C_S1B_included_vs_not.csv、R56_D_S1B_ipw_rates.csv"""
import os, re, numpy as np, pandas as pd
SRC = '05_統计代码'.replace('統', '统') + '/可移植版/顶刊修订/'; TAB = '11_顶刊修订分析/tables/'
src = open(SRC + '63_baseline_phenotype_analysis.py', encoding='utf-8').read().split('# ---------- 1. availability')[0]
g = {'__name__': 'r63'}; exec(src, g)
IA, I = g['I_ALL'], g['I']
nu = lambda x: x.astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
# E30 修正（2026-10-09）：患者唯一号缺失时原来被转成字符串 'nan' 并合并为一例（配对集因此得 836 例）；改为与 23/27a 号脚本一致，缺失时用登记号
pkey = lambda u, r: nu(u).where(nu(u).ne('nan'), 'R' + nu(r))
arch = IA[IA.grade_source == 'TTE文本解析'].copy(); arch['uid'] = pkey(arch['患者唯一号'], arch['登记号'])
excl_arch = arch[(arch.excl_valve_surgery | arch.excl_team_review) & ~arch.reinstated]
# registry branch (same rules as scripts 19/23)
import openpyxl
raw = '09_原始数据_含直接标识符/出院时间2020-1-2024-9总表改2.xlsx'
ws = openpyxl.load_workbook(raw)['Sheet1']
fill = lambda c: (c.fill.fgColor.rgb if (c.fill is not None and c.fill.fill_type is not None and c.fill.fgColor.type == 'rgb') else 'none')
R = pd.read_excel(raw); col = lambda p: R.iloc[:, p - 1]
f3 = np.array([fill(ws.cell(r, 3)) for r in range(2, len(R) + 2)]); f4 = np.array([fill(ws.cell(r, 4)) for r in range(2, len(R) + 2)])
mr = pd.to_numeric(col(91), errors='coerce'); tr = pd.to_numeric(col(92), errors='coerce'); af_rec = col(10).astype(str).str.strip().ne('nan')
grp = pd.to_numeric(col(6), errors='coerce'); uid = pkey(col(7), col(4))
nograde = mr.isna() & tr.isna()
reg = ~((f3 == 'FFFF0000') | (nograde & ~af_rec)) & ~(f4 == 'FFFFFF00') & ((mr >= 1) | (tr >= 1))
vr = (f4 == 'FFFF0000')
assert int(reg.sum()) == 1479 and int((reg & vr).sum()) == 8 and int((reg & grp.isin([1, 2, 3])).sum()) == 844
trunk = set(arch.uid) | set(uid[reg])
ex_pat = set(excl_arch.uid) | set(uid[reg & vr])
study = trunk - ex_pat
eligible = reg & ~vr & ~uid.isin(set(excl_arch.uid))
paired = eligible & grp.isin([1, 2, 3]); nof = eligible & ~grp.isin([1, 2, 3])
assert uid[paired].nunique() == 837, uid[paired].nunique()   # 与 27a 号聚类键一致；若 76 号核查判定应为 844，需先改患者键
base = I[I.text]
flow = [('study population before exclusion: patients', len(trunk)),
        ('  of whom index admission with archived report (baseline set source)', len(set(arch.uid))),
        ('  registry admissions with MR/TR >=1', int(reg.sum())),
        ('excluded patients (uniform criteria)', len(ex_pat)),
        ('study population: patients', len(study)),
        ('baseline phenotype analysis set: patients', len(base)),
        ('  PAF / PeAF', f"{int((base.af_type == 'PAF').sum())} / {int((base.af_type == 'PeAF').sum())}"),
        ('eligible registry admissions with MR/TR >=1', int(eligible.sum())), ('  patients', uid[eligible].nunique()),
        ('no 1-month TTE: admissions', int(nof.sum())), ('  patients', uid[nof].nunique()),
        ('paired response set: procedures', int(paired.sum())), ('  patients', uid[paired].nunique()),
        ('patients in both analysis sets', len(set(base.uid) & set(uid[paired]))),
        ('study-population patients in neither analysis set', len(study - set(base.uid) - set(uid[paired])))]
pd.DataFrame(flow, columns=['step', 'n']).to_csv(TAB + 'R56_A_unified_flow.csv', index=False)
print(pd.DataFrame(flow).to_string(index=False))
# exclusion categories (archive-side team review + automated valve-surgery rule)
REV = g['REV']; REV['uid'] = nu(REV['患者唯一号'])
cats = []
for _, r in excl_arch.iterrows():
    if r.excl_valve_surgery: cats.append('Prior or concurrent valve surgery')
    else:
        why = REV.loc[REV.uid == r.uid, '初筛理由']
        cats.append(why.iloc[0] if len(why) else 'unknown')
cc = pd.Series(cats).value_counts()
extra_reg_vr = len(set(uid[reg & vr]) - set(excl_arch.uid))
cc.loc['Prior or concurrent valve surgery (registry only)'] = extra_reg_vr
cc.to_csv(TAB + 'R56_B_exclusion_categories.csv'); print(cc)
# ---------------- Table S1B recomputed (same definitions as script 19)
D = pd.DataFrame({'grp': grp, 'pid': uid, 'mr_raw': mr, 'tr_raw': tr, 'mr_b': pd.to_numeric(col(93), errors='coerce'), 'tr_b': pd.to_numeric(col(95), errors='coerce'),
                  'mr_1': pd.to_numeric(col(157), errors='coerce'), 'tr_1': pd.to_numeric(col(158), errors='coerce'), 'af': col(10).astype(str).str.strip(),
                  'sex': col(12).astype(str).str.strip(), 'age': pd.to_numeric(col(13), errors='coerce'), 'LA': pd.to_numeric(col(105), errors='coerce'),
                  'RV': pd.to_numeric(col(106), errors='coerce'), 'RA': pd.to_numeric(col(107), errors='coerce'), 'LVEF': pd.to_numeric(col(127), errors='coerce')})
E = D[eligible].copy(); E['included'] = E.grp.isin([1, 2, 3]).astype(int)
rec_mr = lambda x: np.select([x.isna(), x == 0, x.isin([1, 2]), x.isin([3, 4]), x >= 5], [np.nan, 0, 1, 2, 3], np.nan)
rec_tr = lambda x: np.select([x.isna(), x == 0, x.isin([1, 2]), x.isin([3, 4]), x.isin([5, 6]), x >= 7], [np.nan, 0, 1, 2, 3, 4], np.nan)
E['MRg'] = rec_mr(E.mr_raw); E['TRg'] = rec_tr(E.tr_raw)
E['pheno'] = np.select([(E.MRg >= 1) & (E.TRg >= 1), E.MRg >= 1, E.TRg >= 1], ['C', 'A', 'B'], 'NA')
E['female'] = (E.sex == '女').astype(float); E.loc[~E.sex.isin(['男', '女']), 'female'] = np.nan
afu = E.af.str.upper().str.strip(); E['persistent'] = np.select([afu.isin(['PEAF', '2']), afu.isin(['PAF', '1'])], [1.0, 0.0], np.nan)
for v, lo_, hi_ in [('LA', 20, 80), ('RA', 20, 90), ('RV', 10, 60), ('LVEF', 10, 85)]: E.loc[(E[v] < lo_) | (E[v] > hi_), v] = np.nan
E['MR_ge_mod'] = (E.MRg >= 2).astype(float); E['TR_ge_mod'] = (E.TRg >= 2).astype(float)
def smd(x1, x0, b):
    x1, x0 = x1.dropna(), x0.dropna(); p1, p0 = x1.mean(), x0.mean()
    s = np.sqrt((p1 * (1 - p1) + p0 * (1 - p0)) / 2) if b else np.sqrt((x1.var() + x0.var()) / 2); return (p1 - p0) / s
inc, exc = E[E.included == 1], E[E.included == 0]; t = []
for v, lab, b in [('age', 'Age, y', 0), ('female', 'Female sex', 1), ('persistent', 'Persistent AF', 1), ('LA', 'LA diameter, mm', 0), ('RA', 'RA diameter, mm', 0),
                  ('RV', 'RV diameter, mm', 0), ('LVEF', 'LVEF, %', 0), ('MR_ge_mod', 'MR grade >=2', 1), ('TR_ge_mod', 'TR grade >=2', 1)]:
    f = (lambda x: f'{100 * x.mean():.1f}%') if b else (lambda x: f'{x.mean():.1f} ± {x.std():.1f}')
    t.append(dict(variable=lab, included=f(inc[v].dropna()), not_included=f(exc[v].dropna()), SMD=round(smd(inc[v], exc[v], b), 2), n_not_included=int(exc[v].notna().sum())))
for ph, lab in (('A', 'isolated MR'), ('B', 'isolated TR'), ('C', 'combined MR/TR')):
    x1, x0 = (inc.pheno == ph).astype(float), (exc.pheno == ph).astype(float)
    t.append(dict(variable=lab, included=f'{100 * x1.mean():.1f}%', not_included=f'{100 * x0.mean():.1f}%', SMD=round(smd(x1, x0, 1), 2), n_not_included=len(exc)))
T7 = pd.DataFrame(t); T7.to_csv(TAB + 'R56_C_S1B_included_vs_not.csv', index=False); print(T7.to_string(index=False))
print('AF type missing among not included:', int(exc.persistent.isna().sum()), '| eligible', len(E), 'patients', E.pid.nunique())
# IPW for follow-up (as script 19)
def logistic(X, y, it=50):
    b = np.zeros(X.shape[1])
    for _ in range(it):
        p = 1 / (1 + np.exp(-X @ b)); W = p * (1 - p); st = np.linalg.solve(X.T @ (X * W[:, None]), X.T @ (y - p)); b += st
        if np.abs(st).max() < 1e-9: break
    return b
Z = E.copy(); covs = ['age', 'female', 'persistent', 'LA', 'LVEF', 'RA']
for v in covs: Z[v + '_miss'] = Z[v].isna().astype(float); Z[v] = Z[v].fillna(Z[v].median())
Z['phB'] = (Z.pheno == 'B').astype(float); Z['phC'] = (Z.pheno == 'C').astype(float); Z['MRg2'] = (Z.MRg >= 2).astype(float); Z['TRg2'] = (Z.TRg >= 2).astype(float)
Xc = covs + [v + '_miss' for v in covs if Z[v + '_miss'].sum() > 0] + ['phB', 'phC', 'MRg2', 'TRg2']
X = np.column_stack([np.ones(len(Z))] + [((Z[c] - Z[c].mean()) / (Z[c].std() or 1)).values for c in Xc]); y = Z.included.values.astype(float)
ps = 1 / (1 + np.exp(-X @ logistic(X, y)))
r = np.argsort(np.argsort(ps)); cstat = (r[y == 1].mean() - (y.sum() - 1) / 2) / (len(y) - y.sum()) / 1 if False else None
o = ps.argsort(); ranks = np.empty(len(ps)); ranks[o] = np.arange(1, len(ps) + 1); cstat = (ranks[y == 1].sum() - y.sum() * (y.sum() + 1) / 2) / (y.sum() * (len(y) - y.sum()))
w = np.where(y == 1, 1 / ps, 0); lo_, hi_ = np.percentile(w[y == 1], [1, 99]); Z['w'] = np.clip(w, lo_, hi_) * y.mean()
rng = np.random.default_rng(20260924); t8 = []
for k, (gg, b_, a_) in {'Isolated MR': (1, 'mr_b', 'mr_1'), 'Isolated TR': (2, 'tr_b', 'tr_1'), 'Combined disease: MR': (3, 'mr_b', 'mr_1'), 'Combined disease: TR': (3, 'tr_b', 'tr_1')}.items():
    x = Z[(Z.included == 1) & (Z.grp == gg)]; ev = (x[a_] < x[b_]).values.astype(float); ww = x.w.values
    bs = [(ev[s] * ww[s]).sum() / ww[s].sum() for s in (rng.integers(0, len(x), len(x)) for _ in range(2000))]
    t8.append(dict(series=k, n=len(x), crude=round(100 * ev.mean(), 1), weighted=round(100 * (ev * ww).sum() / ww.sum(), 1), lo=round(np.percentile(bs, 2.5) * 100, 1), hi=round(np.percentile(bs, 97.5) * 100, 1)))
T8 = pd.DataFrame(t8); T8['c_stat'] = round(cstat, 2); T8['weight_range'] = f"{Z.loc[Z.included == 1, 'w'].min():.2f}–{Z.loc[Z.included == 1, 'w'].max():.2f}"
T8.to_csv(TAB + 'R56_D_S1B_ipw_rates.csv', index=False); print(T8.to_string(index=False))
miss = {v: int(E[v].isna().sum()) for v in ('persistent', 'LA', 'RA', 'LVEF')}; print('missing in eligible:', miss)
